from pathlib import Path
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from schema_diagrams import generate_schema_diagrams

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get('DBMS_REPORT_OUTPUT', ROOT / 'report'))
ASSET = OUT / 'assets'
ASSET.mkdir(exist_ok=True)

def font(size=12, bold=False, italic=False):
    f = ImageFont.load_default()
    return f

def box(draw, xy, text, fill='#F1F2F3', outline='#5B6670'):
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
    x1,y1,x2,y2=xy
    draw.text(((x1+x2)//2, (y1+y2)//2), text, anchor='mm', fill='#132238', font=font())

def make_er():
    img=Image.new('RGB',(1200,900),'white'); d=ImageDraw.Draw(img)
    box(d,(40,70,270,260),'PROFILES'); box(d,(480,70,760,260),'ACCOUNTS'); box(d,(900,70,1160,260),'BRANCHES')
    box(d,(170,430,500,640),'TRANSACTIONS', '#F2F8F4','#16805A'); box(d,(700,430,1060,640),'BENEFICIARIES','#FFF7E6','#C67C00')
    box(d,(390,700,810,860),'RECOVERY DEMO EVENTS','#FFF1E8','#C67C00')
    for a,b in [((270,165),(480,165)),((760,165),(900,165)),((400,260),(330,430)),((650,260),(850,430)),((230,260),(230,430)),((390,260),(480,700))]: d.line((a,b),fill='#64748B',width=4)
    d.text((600,25),'Entity Relationship Diagram',anchor='mm',fill='#132238',font=font())
    img.save(ASSET/'er_diagram_final.png')

def make_arch():
    img=Image.new('RGB',(1200,650),'white'); d=ImageDraw.Draw(img)
    box(d,(60,240,300,400),'Customer / Staff / Admin'); box(d,(480,240,720,400),'Browser UI'); box(d,(900,240,1140,400),'Supabase')
    for x1,x2 in [(300,480),(720,900)]: d.line((x1,320,x2,320),fill='#5B6670',width=5); d.polygon([(x2,320),(x2-18,310),(x2-18,330)],fill='#5B6670')
    d.text((600,90),'System Architecture and Access Flow',anchor='mm',fill='#132238',font=font())
    for x,text in [(180,'Authentication'),(600,'Database requests'),(1020,'PostgreSQL + RLS')]: d.text((x,500),text,anchor='mm',fill='#5B6670',font=font())
    img.save(ASSET/'architecture.png')

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_border(cell, color='D9D9D9'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),color)

def add_table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; set_cell_shading(c,'E6E8EA'); set_cell_border(c,'AEB4BA'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(32,42,51); r.font.bold=True
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); set_cell_border(cells[i],'D2D6DA'); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: set_cell_shading(cells[i],'F6F7F8')
    return t

def p(doc, text='', bold=False, italic=False, align=None):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(6); para.paragraph_format.line_spacing=1.5
    if align is not None: para.alignment=align
    run=para.add_run(text); run.bold=bold; run.italic=italic; run.font.name='Times New Roman'; run.font.size=Pt(12); run.font.color.rgb=RGBColor(0,0,0)
    return para

def heading(doc, text, level=1):
    para=doc.add_paragraph(style=f'Heading {level}'); para.paragraph_format.space_before=Pt(12); para.paragraph_format.space_after=Pt(6)
    run=para.add_run(text); run.font.name='Times New Roman'; run.font.color.rgb=RGBColor(0,0,0); run.bold=True; run.font.size=Pt(14 if level==1 else 12)
    return para

def page_break(doc): doc.add_page_break()

def add_page_num(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run('Page '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

def build():
    generate_schema_diagrams(ASSET)
    make_arch()
    doc=Document(); sec=doc.sections[0]
    sec.left_margin=Inches(1.5); sec.right_margin=Inches(1); sec.top_margin=Inches(1); sec.bottom_margin=Inches(1)
    styles=doc.styles; styles['Normal'].font.name='Times New Roman'; styles['Normal'].font.size=Pt(12)
    for s in sec.footer.paragraphs: add_page_num(s)
    # Cover
    p(doc,'Vivekanand Education Society\'s Institute of Technology',True,align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc,'Department of Computer Engineering',True,align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc,'Academic Year 2026-27',align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc,'',align=WD_ALIGN_PARAGRAPH.CENTER); p(doc,'DBMS Micro Project Report',True,align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc,'Digital Banking System Database',True,align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc,'A Supabase PostgreSQL implementation with a role-aware web frontend',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc,'',align=WD_ALIGN_PARAGRAPH.CENTER); p(doc,'Course: Database Management System (DBMS)',align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)
    heading(doc,'Index'); add_table(doc,['No.','Section'],[(1,'Introduction and Background'),(2,'Problem Statement'),(3,'Objectives and Scope'),(4,'Requirements and Data Requirements'),(5,'ER and Relational Schema Diagrams'),(6,'Flow Diagram and System Architecture'),(7,'Implementation and Database Operations'),(8,'Results and Working System Screenshots'),(9,'Conclusion and Future Scope')])
    p(doc,'This report is prepared according to the prescribed DBMS micro-project format. The implementation uses Supabase as the hosted PostgreSQL database and authentication platform.')
    page_break(doc)
    heading(doc,'1. Introduction and Background')
    p(doc,'Digital banking lets customers check accounts and make payments through a web or mobile interface. Behind each screen, the bank must keep account balances, customer records, branches, and transaction history accurate and protected. Transfers also need to remain consistent if an operation fails partway through.')
    p(doc,'This project builds a small digital banking database with Supabase and PostgreSQL. It brings together the database, authentication, access rules, and a web interface so the main banking workflows can be demonstrated with real records.')
    p(doc,'A signed-in user can access only the records allowed by their role. Transfers run as a single database operation: the debit, credit, and transaction record either all succeed or all roll back.')
    page_break(doc)
    heading(doc,'2. Problem Statement')
    p(doc,'A banking interface can look convincing while still relying on unrelated files or temporary frontend data. That makes it hard to keep balances correct, connect transactions to the right accounts, or control who can see sensitive information.')
    p(doc,'This project addresses that problem with a relational database for customers, branches, accounts, beneficiaries, and transactions. It pairs those records with role-based access for customers, staff, and administrators, and presents the main operations through a simple web interface.')
    page_break(doc)
    heading(doc,'3. Objectives and Scope')
    heading(doc,'3.1 Objectives',2)
    for x in ['Design a relational schema for core digital banking entities.','Apply primary keys, foreign keys, unique constraints, checks, indexes, and views.','Implement authentication-aware access using Supabase Row Level Security.','Provide an atomic money-transfer function using PostgreSQL transactions.','Build a basic frontend for dashboard viewing, balance lookup, transfers, and activity history.','Demonstrate SQL selection, joins, aggregation, updates, and stored procedure calls.'] : p(doc,'• '+x)
    heading(doc,'3.2 Scope',2)
    p(doc,'The project covers customer accounts, savings and current account types, branch association, beneficiaries, deposits, withdrawals, transfers, account status, and role-based visibility. It is an educational prototype and does not implement production banking integrations such as payment gateways, regulatory reporting, KYC document storage, OTP delivery, interest calculation, or a complete fraud-monitoring system.')
    page_break(doc)
    heading(doc,'4. Requirements and Data Requirements')
    heading(doc,'4.1 Functional Requirements',2)
    add_table(doc,['Requirement','Description'],[('Authentication','Users sign in through Supabase Auth.'),('Dashboard','Customers view their own accounts and balances.'),('Transfer','A signed-in customer can transfer to an active destination account.'),('History','Users view financial transactions and durable interrupted-transfer demo events.'),('Beneficiaries','Customers maintain their beneficiaries; admins review and remove beneficiary records.'),('Staff review','Staff can review masked customer accounts and apply limited status actions.'),('Administration','Admins oversee system metrics, account states, provisioning, and beneficiary records.')])
    heading(doc,'4.2 Non-Functional Requirements',2)
    p(doc,'The system protects customer records with Row Level Security, keeps related data consistent with constraints and foreign keys, and provides a responsive interface for common tasks. Separate SQL scripts make setup and database demonstrations easier to follow. Indexes support account lookups and transaction history queries.')
    heading(doc,'4.3 Data Requirements',2)
    add_table(doc,['Entity','Important attributes'],[('Branches','branch code, name, address, city'),('Profiles','auth user ID, name, email, role, branch'),('Accounts','account number, owner, type, balance, status'),('Beneficiaries','owner, nickname, account number, bank, IFSC'),('Transactions','source, destination, amount, type, status, initiator, time'),('Recovery demo events','customer, source account, masked destination, amount, rolled-back status, detail, time')])
    diagram_sec=doc.add_section(WD_SECTION.NEW_PAGE)
    diagram_sec.orientation=WD_ORIENT.LANDSCAPE
    diagram_sec.page_width=Inches(11); diagram_sec.page_height=Inches(8.5)
    diagram_sec.left_margin=Inches(.45); diagram_sec.right_margin=Inches(.45)
    diagram_sec.top_margin=Inches(.4); diagram_sec.bottom_margin=Inches(.4)
    # Give diagrams their own landscape pages so table labels remain legible in print.
    heading(doc,'5. ER and Relational Schema Diagrams')
    doc.add_picture(str(ASSET/'digital_banking_chen_er.png'), width=Inches(9.9))
    p(doc,'Figure 1. Chen ER diagram. Underlined attributes are primary keys.',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)
    doc.add_picture(str(ASSET/'digital_banking_relational_schema.png'), width=Inches(9.9))
    p(doc,'Figure 2. Relational schema. Each FK field names its referenced key after the arrow.',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    portrait_sec=doc.add_section(WD_SECTION.NEW_PAGE)
    portrait_sec.orientation=WD_ORIENT.PORTRAIT
    portrait_sec.page_width=Inches(8.5); portrait_sec.page_height=Inches(11)
    portrait_sec.left_margin=Inches(1.5); portrait_sec.right_margin=Inches(1)
    portrait_sec.top_margin=Inches(1); portrait_sec.bottom_margin=Inches(1)
    heading(doc,'6. Flow Diagram and System Architecture')
    p(doc,'The frontend uses the publishable Supabase key. Authentication establishes the signed-in user. RLS policies then limit table access, while the transfer function validates ownership, balance, status, and destination account before updating records.')
    doc.add_picture(str(ASSET/'architecture.png'), width=Inches(6.2)); p(doc,'Figure 3: High-level architecture and access flow.',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    heading(doc,'6.1 Access Control Flow',2)
    p(doc,'Customer requests are filtered by auth.uid(). Staff and admin access is determined by the role stored in profiles. The role helper function is security definer and uses a restricted search path. The browser never receives the service-role key; all frontend access is intended to occur through the publishable key and database policies.')
    page_break(doc)
    heading(doc,'7. Implementation and Database Operations')
    heading(doc,'7.1 Schema Implementation',2)
    p(doc,'The database defines enum types for user roles, account types, account statuses, transaction types, and transaction statuses. Account states include active, frozen, and closed. These states are represented in seeded customer records so the interface can demonstrate how account access changes with status. Accounts use a non-negative balance check. Transactions require a positive amount and at least one related account.')
    heading(doc,'7.2 Database Operations Demonstrated',2)
    add_table(doc,['Operation','Example used in project','DBMS concept'],[('SELECT','Account summary and transaction history','Filtering and ordering'),('JOIN','Branch-wise balances','Relational association'),('GROUP BY','Daily transfer totals','Aggregation'),('INSERT','New beneficiary or transaction','Data creation'),('UPDATE','Profile phone or account status','Data modification'),('RPC function','Atomic transfer','Stored procedure and consistency'),('RLS policy','Own rows versus staff rows','Authorization')])
    heading(doc,'7.3 Atomic Transfer',2)
    p(doc,'The transfer_money function validates a positive amount and distinct accounts. It locks the source and destination rows in a consistent order, deducts the amount only from an active account owned by the signed-in user and with sufficient balance, credits an active destination, and inserts a completed transaction record. PostgreSQL executes the RPC statement within a transaction. If the client disconnects, the statement is cancelled, or a database error occurs before commit, the database rolls back every change from that statement.')
    heading(doc,'7.3.1 Recovery Demonstration',2)
    p(doc,'The demo_interrupted_transfer function performs the simulated debit and pending transaction insert inside a nested database block, then deliberately raises a controlled exception. PostgreSQL rolls back that block, restoring the source balance and removing the pending financial ledger row. The outer function then writes a separate recovery_demo_events audit entry containing the amount, source account, masked destination, timestamp, and rollback detail. The customer activity history displays this event as “Interrupted transfer - rolled back.” The event records the demonstration without representing a completed or pending financial transfer.')
    page_break(doc)
    heading(doc,'7.4 Security and Permissions')
    add_table(doc,['Role','Workspace and authority'],[('Customer','Personal balances, account states, saved beneficiaries, transaction history, and transfers from active accounts.'),('Staff','Masked account review; may freeze an active account or reactivate a frozen account. Cannot change balances, assign roles, or close accounts.'),('Admin','System metrics, all account-state controls, account provisioning for existing customers, and beneficiary review and removal.')])
    p(doc,'Row Level Security is enabled on the core tables. Direct account and ledger writes are revoked from browser roles; controlled security-definer functions perform transfers and management actions after checking the caller’s role. Profile updates are limited to contact fields, so a user cannot promote themselves by changing their role. The account summary view returns masked account numbers and does not include customer names or email addresses.')
    heading(doc,'7.5 Frontend Implementation',2)
    p(doc,'The frontend is Supabase-only and requires sign-in through Supabase Auth; it has no demo preview or fabricated role data. It provides a customer banking workspace, an operations workspace for staff, and a system control center for administrators. Staff controls are limited to freeze and reactivate actions. Administrators manage account states, provision accounts, and review or remove beneficiary records; the former role-assignment panel is not shown. Staff and administrator account lists use masked account numbers. Customer history combines financial transactions with separately stored recovery-demo audit events.')
    page_break(doc)
    heading(doc,'7.6 SQL Command Samples')
    p(doc,'The following commands are included in the project so the database can be demonstrated directly in the Supabase SQL Editor.')
    add_table(doc,['Command category','Demonstration'],[('DDL','CREATE TABLE, CREATE TYPE, CREATE INDEX, CREATE VIEW'),('DML','INSERT branch, UPDATE profile, and stored transaction insert'),('DQL','SELECT account summary, JOIN branches, GROUP BY transfer totals'),('Security','ALTER TABLE ENABLE ROW LEVEL SECURITY and CREATE POLICY'),('Procedural','SELECT transfer_money(...) for an atomic transfer')])
    p(doc,'A useful demonstration sequence is to run the account summary query as a customer, run the same query with a different user, and then run the branch-wise aggregation as staff. The result difference shows that database permissions are part of the system behavior.')
    heading(doc,'7.7 Data Dictionary',2)
    add_table(doc,['Table','Key fields','Purpose'],[('branches','branch_id, branch_code','Stores physical branch information.'),('profiles','user_id, role, branch_id','Connects Auth identity to banking profile and role.'),('accounts','account_id, account_number, balance','Stores customer accounts and current balance.'),('beneficiaries','beneficiary_id, user_id','Stores saved transfer destinations.'),('transactions','transaction_id, amount, status','Stores financial activity and audit history.'),('recovery_demo_events','event_id, user_id, from_account_id, amount','Records a rolled-back recovery demo separately from the financial ledger.')])
    page_break(doc)
    heading(doc,'8. Results and Working System Screenshots')
    p(doc,'The current demonstration data includes customer profiles, accounts in active, frozen, and closed states, saved beneficiaries, and sample transactions. Additional customer records can be loaded after their Supabase Auth users are created. The following evidence can be collected during a live demonstration:')
    for x in ['Supabase tables showing branches, profiles, accounts, beneficiaries, transactions, and recovery_demo_events.','SQL Editor showing installed schema, RLS policies, sample records, and demonstration queries.','Auth users linked to customer, staff, and admin profiles.','Customer dashboard with account states, saved beneficiaries, and transaction history.','Interrupted-transfer entry marked rolled back in the activity panel, alongside unchanged balances and no pending financial ledger row.','Staff account review showing masked numbers and limited status controls.','Admin console showing system metrics, beneficiary management, account lifecycle controls, and account provisioning.','Successful transfer with updated balances and a completed ledger record.'] : p(doc,'• '+x)
    heading(doc,'8.1 Expected Results',2)
    add_table(doc,['Test case','Expected result'],[('Customer reads accounts','Only accounts owned by signed-in customer are returned.'),('Customer reads another profile','Request is rejected or returns no row.'),('Valid transfer','Source decreases, destination increases, transaction is inserted.'),('Interrupted transfer demo','Debit and pending ledger row roll back; a separate rolled-back event appears in activity history.'),('Insufficient balance','Function raises an error and balances remain unchanged.'),('Staff status action','Staff can freeze or reactivate; closing an account is rejected.'),('Admin beneficiary controls','Admin can review beneficiary owners and masked destination numbers, then remove a selected record.'),('Admin account control','Admin can manage account states and provision an account.'),('Masked staff/admin view','Account summary displays only the final four account-number digits.'),('Invalid amount','Database check/function rejects zero or negative amount.')])
    page_break(doc)
    heading(doc,'8.2 Test Plan and Demonstration Evidence')
    p(doc,'Testing combines database-level checks with interface-level checks. Each test should be run after the Supabase project has been configured with Auth users and sample accounts.')
    add_table(doc,['Test ID','Action','Expected evidence'],[('T01','Sign in as customer','Personal dashboard and role badge appear.'),('T02','Load accounts','Only own accounts appear with their current state.'),('T03','Run interrupted transfer','The financial debit and pending row roll back; a separate recovery event appears in the activity list.'),('T04','Complete a valid transfer','Balances and recent activity update.'),('T05','Attempt an over-balance transfer','The function raises an error; no partial update occurs.'),('T06','Sign in as staff','Masked review table and limited status actions appear.'),('T07','Sign in as admin','Metrics, beneficiary review/removal, account status, and account provisioning controls appear.'),('T08','Inspect RLS','Unauthorized rows and direct writes are denied.')])
    p(doc,'For the final viva, the demonstrator should show the UI action and then immediately show the corresponding database row or balance change in the SQL Editor. This connects the visible frontend behavior to the underlying DBMS operations.')
    heading(doc,'8.3 Limitations of the Prototype',2)
    p(doc,'This project is an academic prototype, not a production banking service. A real deployment would need a broader security review, operational monitoring, rate limits, audit-retention rules, compliance controls, and testing with external banking services.')
    page_break(doc)
    heading(doc,'Appendix A: Project Demonstration Procedure')
    for x in ['Use the connected Supabase project and confirm the schema and RLS policies are installed.','Create the required Auth users and link each Auth UUID to the correct banking profile.','Load the additional demo customers and sample financial data using demo_customers.sql after replacing its UUID placeholders.','Start the Supabase-only frontend and sign in with a project Auth user.','Run the interrupted-transfer demonstration; verify the recovery event in customer activity and confirm balances and the financial ledger are unchanged.','Perform a valid transfer and verify both balances and the completed ledger record.','Sign in as staff to demonstrate masked accounts and limited account-status duties.','Sign in as admin to demonstrate system metrics, beneficiary review/removal, account lifecycle control, and provisioning.'] : p(doc,'• '+x)
    heading(doc,'Appendix B: Submission Checklist',2)
    add_table(doc,['Item','Status'],[('Supabase schema script','Included'),('Seed and setup guidance','Included'),('RLS policies','Included'),('Demonstration queries','Included'),('Frontend scaffold','Included'),('ER diagram','Included in report'),('System architecture diagram','Included in report'),('Live screenshots','To capture after Supabase setup'),('Final report formatting','Applied to DOCX draft')])
    page_break(doc)
    heading(doc,'9. Conclusion and Future Scope')
    p(doc,'The project brings together a relational schema, Supabase authentication, Row Level Security, and a role-aware interface to demonstrate common digital banking workflows. Database constraints and an atomic transfer function help keep account data consistent, while the diagrams and SQL examples make the design easier to inspect.')
    p(doc,'The main lesson is that the database must enforce ownership and permissions itself. The frontend helps users work with their accounts, while database policies determine which records they can actually access.')
    heading(doc,'9.1 Future Scope',2)
    for x in ['Add OTP or multi-factor authentication for sensitive actions.','Add beneficiary verification and transfer limits.','Introduce an audit log for staff and admin actions.','Add scheduled payments, bill payments, and notifications.','Add fraud rules and unusual-activity alerts.','Deploy the frontend through a secure hosting platform and connect it to a production Supabase project.'] : p(doc,'• '+x)
    heading(doc,'References')
    p(doc,'Supabase Documentation, PostgreSQL Documentation, and the supplied DBMS Micro Project Report format.')
    output=OUT/'Digital_Banking_System_DBMS_Report.docx'; doc.save(output); print(output)

if __name__=='__main__': build()
